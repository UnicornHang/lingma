"""[P3.4] DOCX 导出 —— python-docx 实现,显式配置 CJK eastAsia 字体。"""
from __future__ import annotations

from typing import Sequence

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.oxml.ns import qn

from app.models.work import Genre, Work
from app.schemas.export import ExportOptions
from app.services.export.base import BaseExporter
from app.services.export.volume_grouping import Volume

_GENRE_LABEL = {
    Genre.FANTASY: "奇幻",
    Genre.URBAN: "都市",
    Genre.ROMANCE: "言情",
    Genre.HISTORICAL: "历史",
    Genre.SCI_FI: "科幻",
    Genre.MYSTERY: "悬疑",
    Genre.OTHER: "其他",
}


def _apply_cjk_font(style, font_name: str) -> None:
    """设置 docx style 的 CJK 字体(同时设 latin + eastAsia)。

    python-docx 默认只设 latin 字体,这对英文生效;中文需要单独写 eastAsia rPr。
    """
    style.font.name = font_name  # 西文
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def _split_paragraphs(text: str) -> list[str]:
    """按双换行(或单换行)拆段;保留段落顺序,空段丢弃。"""
    if not text:
        return []
    parts = [p.strip() for p in text.replace("\r\n", "\n").split("\n\n")]
    out: list[str] = []
    for p in parts:
        if not p:
            continue
        # 段内仍可能含单换行(如对话换行)—— 拆成多段而非合并
        for line in p.split("\n"):
            if line.strip():
                out.append(line.strip())
    return out


class DocxExporter(BaseExporter):
    format = "docx"

    @property
    def media_type(self) -> str:
        return (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def render(
        self,
        work: Work,
        volumes: Sequence[Volume],
        options: ExportOptions,
    ) -> bytes:
        doc = Document()

        # 设置全局 Normal 样式 CJK 字体
        normal_style = doc.styles["Normal"]
        _apply_cjk_font(normal_style, options.cjk_font_name)

        # 标题样式(Heading 1 / 2 / 3)同样设 CJK
        for level in (1, 2, 3):
            try:
                h = doc.styles[f"Heading {level}"]
            except KeyError:
                continue
            _apply_cjk_font(h, options.cjk_font_name)

        # ===== 封面 =====
        cover_title = doc.add_paragraph()
        cover_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = cover_title.add_run(work.title)
        run.bold = True
        run.font.size = Pt(28)
        _apply_cjk_font(cover_title.style, options.cjk_font_name)

        if work.logline:
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.add_run(work.logline)
            run.font.size = Pt(14)

        genre_label = _GENRE_LABEL.get(work.genre, str(work.genre.value))
        meta = doc.add_paragraph()
        meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        meta.add_run(f"\n\n[{genre_label}]   总字数 {work.word_count:,}").font.size = Pt(11)

        # 封面页结束 → 强制分页
        doc.add_page_break()

        # ===== 目录占位 =====
        doc.add_heading("目录", level=1)
        for vol in volumes:
            if len(volumes) > 1 or vol.title != "正文":
                doc.add_paragraph().add_run(f"  {vol.title}").bold = True
            for ch in vol.chapters:
                p = doc.add_paragraph()
                p.add_run(f"    {ch.title}").italic = True
        doc.add_page_break()

        # ===== 正文 =====
        for vol in volumes:
            # 单卷且标题是默认"正文"时,不显示卷标题(避免冗余)
            show_volume_heading = (
                len(volumes) > 1 or vol.title not in ("正文", "未分类")
            )
            if show_volume_heading:
                doc.add_heading(vol.title, level=1)

            for ch in vol.chapters:
                doc.add_heading(ch.title, level=2)
                for para_text in _split_paragraphs(ch.plain_content):
                    doc.add_paragraph(para_text)

        # python-docx 的 save 默认走物理文件;用 BytesIO 拿字节
        import io
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
